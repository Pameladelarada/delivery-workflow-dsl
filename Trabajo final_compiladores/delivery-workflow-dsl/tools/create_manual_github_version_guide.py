from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Guia_subir_version_nueva_github.docx"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(table, top=80, start=120, bottom=80, end=120):
    tbl_pr = table._tbl.tblPr
    margins = tbl_pr.first_child_found_in("w:tblCellMar")
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        tbl_pr.append(margins)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(34, 43, 54)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.add_run(text)


def add_step(doc, title, text, command=None):
    doc.add_heading(title, level=2)
    doc.add_paragraph(text)
    if command:
        add_code(doc, command)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.10

for name, size, color, before, after in [
    ("Heading 1", 16, "2E74B5", 16, 8),
    ("Heading 2", 13, "2E74B5", 12, 6),
    ("Heading 3", 12, "1F4D78", 8, 4),
]:
    style = styles[name]
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = title.add_run("Guia manual para subir una version nueva a GitHub")
r.font.name = "Calibri"
r.font.size = Pt(24)
r.font.bold = True
r.font.color.rgb = RGBColor(11, 37, 69)

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(12)
subtitle.add_run("Proyecto: Trabajo final_compiladores").bold = True

doc.add_heading("Objetivo", level=1)
doc.add_paragraph(
    "Esta guia explica como subir manualmente una nueva version del proyecto a GitHub, controlar el historial con commits "
    "y evitar errores comunes como 'fetch first', conflictos durante rebase y fallos en GitHub Actions."
)

doc.add_heading("Antes de subir", level=1)
add_bullet(doc, "Verifica que el proyecto corre localmente con run_web.ps1.")
add_bullet(doc, "No subas .venv, logs, cache ni archivos temporales.")
add_bullet(doc, "Usa mensajes de commit claros para reconocer cada version.")

add_step(
    doc,
    "1. Entrar a la carpeta del proyecto",
    "Abre PowerShell o CMD y ubicate en la carpeta final del trabajo.",
    r'cd "C:\Usil\6to_ciclo_Compiladores\Trabajo final_compiladores"',
)

add_step(
    doc,
    "2. Revisar el estado actual",
    "Este comando muestra que archivos cambiaron y si estas conectado a origin/main.",
    "git status",
)

add_step(
    doc,
    "3. Ver los cambios antes de guardarlos",
    "Sirve para revisar que realmente subiras lo correcto.",
    "git diff",
)

add_step(
    doc,
    "4. Agregar los archivos al commit",
    "Agrega todos los cambios del proyecto. Si quieres agregar solo un archivo, reemplaza el punto por el nombre del archivo.",
    "git add .",
)

add_step(
    doc,
    "5. Crear el commit de la nueva version",
    "El commit guarda una foto del proyecto en el historial local.",
    'git commit -m "Actualiza version del proyecto final"',
)

add_step(
    doc,
    "6. Traer cambios de GitHub antes de subir",
    "Esto evita el error 'fetch first'. Si otra version ya existe en GitHub, Git intentara integrarla con tu version local.",
    "git pull --rebase origin main",
)

doc.add_heading("Si aparece conflicto durante el rebase", level=1)
doc.add_paragraph("Primero revisa que archivos estan en conflicto:")
add_code(doc, "git status\ngit diff --name-only --diff-filter=U")
doc.add_paragraph("Si quieres conservar tu version local en los archivos conflictivos durante rebase, usa:")
add_code(doc, "git checkout --theirs -- nombre_del_archivo\ngit add nombre_del_archivo\ngit rebase --continue")
doc.add_paragraph("Si quieres conservar la version que ya estaba en GitHub durante rebase, usa:")
add_code(doc, "git checkout --ours -- nombre_del_archivo\ngit add nombre_del_archivo\ngit rebase --continue")
doc.add_paragraph("Si quieres cancelar todo el rebase y volver al estado anterior:")
add_code(doc, "git rebase --abort")

add_step(
    doc,
    "7. Subir la nueva version",
    "Cuando no haya conflictos, sube los commits a GitHub.",
    "git push -u origin main",
)

doc.add_heading("Crear una etiqueta de version", level=1)
doc.add_paragraph("Una etiqueta marca una entrega estable. Por ejemplo, para version 1.1:")
add_code(doc, 'git tag -a v1.1 -m "Version 1.1 del proyecto final"\ngit push origin v1.1')

doc.add_heading("GitHub Actions: por que aparecia error", level=1)
doc.add_paragraph(
    "El workflow anterior intentaba ejecutar ./configure, make, make check y make distcheck. Ese flujo corresponde a otros "
    "proyectos de C/C++, pero este proyecto no usa configure ni Makefile. Por eso Actions fallaba aunque el push estuviera bien."
)
doc.add_paragraph("El workflow correcto para este proyecto debe compilar el archivo compiler/main.cpp y ejecutar un ejemplo DSL:")
add_code(
    doc,
    "mkdir -p bin\n"
    "g++ -std=c++17 -O2 -Wall -Wextra compiler/main.cpp -o bin/delivery_compiler\n"
    "./bin/delivery_compiler examples/pedido_basico.dsl"
)

doc.add_heading("Flujo completo recomendado", level=1)
add_code(
    doc,
    r'cd "C:\Usil\6to_ciclo_Compiladores\Trabajo final_compiladores"' + "\n"
    "git status\n"
    "git add .\n"
    'git commit -m "Actualiza version del proyecto final"\n'
    "git pull --rebase origin main\n"
    "git push -u origin main\n"
    'git tag -a v1.1 -m "Version 1.1 del proyecto final"\n'
    "git push origin v1.1"
)

doc.add_heading("Tabla rapida de comandos", level=1)
table = doc.add_table(rows=1, cols=2)
table.style = "Table Grid"
set_cell_margins(table)
set_table_width(table, [2.35, 4.15])
hdr = table.rows[0].cells
hdr[0].text = "Comando"
hdr[1].text = "Uso"
for cell in hdr:
    shade(cell, "F2F4F7")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

rows = [
    ("git status", "Ver estado del repositorio."),
    ("git add .", "Preparar cambios para commit."),
    ("git commit -m \"mensaje\"", "Guardar una version local."),
    ("git pull --rebase origin main", "Integrar cambios remotos antes de subir."),
    ("git push -u origin main", "Subir cambios a GitHub."),
    ("git tag -a v1.1 -m \"mensaje\"", "Crear etiqueta de version."),
    ("git push origin v1.1", "Subir etiqueta a GitHub."),
]
for command, use in rows:
    cells = table.add_row().cells
    cells[0].text = command
    cells[1].text = use

doc.add_heading("Recomendacion final", level=1)
doc.add_paragraph(
    "Para cada nueva entrega, crea un commit con nombre claro, ejecuta git pull --rebase origin main antes de subir, "
    "resuelve conflictos si aparecen y luego ejecuta git push. Usa tags como v1.0, v1.1 o v2.0 para marcar entregas importantes."
)

doc.save(OUT)
print(OUT)
