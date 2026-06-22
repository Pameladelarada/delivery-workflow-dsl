from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Guia_despliegue_railway.docx"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(table, top=80, start=120, bottom=80, end=120):
    tbl_pr = table._tbl.tblPr
    margins = tbl_pr.first_child_found_in("w:tblCellMar")
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        tbl_pr.append(margins)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(34, 43, 54)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.add_run(text)


def add_step(doc, title, body, command=None):
    doc.add_heading(title, level=2)
    doc.add_paragraph(body)
    if command:
        add_code(doc, command)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.10

for name, size, color, before, after in [
    ("Heading 1", 16, "2E74B5", 16, 8),
    ("Heading 2", 13, "2E74B5", 12, 6),
    ("Heading 3", 12, "1F4D78", 8, 4),
]:
    style = styles[name]
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = title.add_run("Guia para desplegar el proyecto en Railway")
r.font.name = "Calibri"
r.font.size = Pt(24)
r.font.bold = True
r.font.color.rgb = RGBColor(11, 37, 69)

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(12)
subtitle.add_run("Proyecto: Trabajo final_compiladores").bold = True

doc.add_heading("Que se configuro en el proyecto", level=1)
doc.add_paragraph(
    "El proyecto fue preparado para desplegarse en Railway usando Docker. Esto permite compilar el compilador C++ "
    "en Linux y ejecutar la aplicacion Flask con Gunicorn."
)
add_bullet(doc, "web/app.py ahora detecta Windows o Linux: en Windows usa bin/delivery_compiler.exe y en Linux usa bin/delivery_compiler.")
add_bullet(doc, "requirements.txt incluye Flask y Gunicorn.")
add_bullet(doc, "Dockerfile instala g++, instala dependencias Python, compila compiler/main.cpp y arranca Gunicorn.")
add_bullet(doc, "railway.json indica a Railway que debe usar el Dockerfile.")
add_bullet(doc, ".dockerignore evita subir al build carpetas innecesarias como .venv, bin, logs y cache.")

doc.add_heading("Archivos agregados o modificados", level=1)
table = doc.add_table(rows=1, cols=2)
table.style = "Table Grid"
set_cell_margins(table)
set_table_width(table, [2.2, 4.3])
hdr = table.rows[0].cells
hdr[0].text = "Archivo"
hdr[1].text = "Proposito"
for cell in hdr:
    shade(cell, "F2F4F7")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

rows = [
    ("Dockerfile", "Define como Railway construye y ejecuta el proyecto en Linux."),
    ("railway.json", "Indica que el builder sera Dockerfile y define healthcheck."),
    (".dockerignore", "Excluye archivos innecesarios del contenedor."),
    ("requirements.txt", "Agrega gunicorn para ejecucion en produccion."),
    ("web/app.py", "Usa el ejecutable correcto segun el sistema operativo y lee el puerto PORT."),
]
for left, right in rows:
    cells = table.add_row().cells
    cells[0].text = left
    cells[1].text = right

doc.add_heading("Contenido principal del Dockerfile", level=1)
add_code(
    doc,
    "FROM python:3.12-slim\n"
    "WORKDIR /app\n"
    "RUN apt-get update && apt-get install -y --no-install-recommends g++\n"
    "COPY requirements.txt .\n"
    "RUN pip install --no-cache-dir -r requirements.txt\n"
    "COPY . .\n"
    "RUN mkdir -p bin && g++ -std=c++17 -O2 -Wall -Wextra compiler/main.cpp -o bin/delivery_compiler\n"
    'CMD [\"sh\", \"-c\", \"gunicorn --bind 0.0.0.0:${PORT:-5000} web.app:app\"]'
)

doc.add_heading("Subir estos cambios a GitHub", level=1)
add_step(
    doc,
    "1. Entrar a la carpeta del proyecto",
    "Abre PowerShell y entra a la carpeta final.",
    r'cd "C:\Usil\6to_ciclo_Compiladores\Trabajo final_compiladores"',
)
add_step(
    doc,
    "2. Revisar archivos modificados",
    "Confirma que aparecen Dockerfile, railway.json, .dockerignore, requirements.txt y web/app.py.",
    "git status",
)
add_step(
    doc,
    "3. Guardar una nueva version local",
    "Crea un commit con las configuraciones de Railway.",
    'git add .\ngit commit -m "Configura despliegue en Railway con Docker"',
)
add_step(
    doc,
    "4. Integrar con GitHub antes de subir",
    "Evita errores tipo fetch first trayendo cambios remotos antes del push.",
    "git pull --rebase origin main",
)
add_step(
    doc,
    "5. Subir a GitHub",
    "Sube la nueva version para que Railway pueda leerla.",
    "git push -u origin main",
)

doc.add_heading("Crear el proyecto en Railway", level=1)
add_step(
    doc,
    "1. Entrar a Railway",
    "Abre Railway e inicia sesion con tu cuenta de GitHub.",
    "https://railway.com",
)
add_step(
    doc,
    "2. Crear un proyecto nuevo",
    "Selecciona New Project y luego Deploy from GitHub repo.",
)
add_step(
    doc,
    "3. Elegir el repositorio",
    "Selecciona el repositorio delivery-workflow-dsl o el nombre que tengas en GitHub.",
)
add_step(
    doc,
    "4. Confirmar que Railway use Dockerfile",
    "Railway debe detectar el Dockerfile de la raiz. Si pregunta el builder, elige Dockerfile.",
)
add_step(
    doc,
    "5. Esperar el build",
    "En los logs deberias ver instalacion de g++, pip install, compilacion de compiler/main.cpp y arranque con Gunicorn.",
)
add_step(
    doc,
    "6. Generar dominio publico",
    "En Railway abre Settings o Networking y selecciona Generate Domain. Ese sera el enlace publico de la pagina.",
)

doc.add_heading("Que revisar en los logs de Railway", level=1)
add_bullet(doc, "Que aparezca la instalacion de g++ sin errores.")
add_bullet(doc, "Que el comando g++ compile compiler/main.cpp.")
add_bullet(doc, "Que se cree bin/delivery_compiler.")
add_bullet(doc, "Que Gunicorn escuche en 0.0.0.0 y use la variable PORT.")

doc.add_heading("Errores comunes", level=1)
table2 = doc.add_table(rows=1, cols=2)
table2.style = "Table Grid"
set_cell_margins(table2)
set_table_width(table2, [2.4, 4.1])
hdr2 = table2.rows[0].cells
hdr2[0].text = "Error"
hdr2[1].text = "Solucion"
for cell in hdr2:
    shade(cell, "F2F4F7")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

errors = [
    ("No se encontro bin/delivery_compiler", "Revisa que el Dockerfile compile con g++ y que el archivo compiler/main.cpp exista."),
    ("Application failed to respond", "Verifica que Gunicorn use 0.0.0.0:${PORT:-5000}."),
    ("ModuleNotFoundError: flask", "Confirma que requirements.txt este incluido y que el build ejecute pip install."),
    ("Build failed en apt-get", "Vuelve a desplegar. Si persiste, revisa los logs completos del build."),
    ("Fetch first al subir a GitHub", "Ejecuta git pull --rebase origin main antes de git push."),
]
for err, sol in errors:
    cells = table2.add_row().cells
    cells[0].text = err
    cells[1].text = sol

doc.add_heading("Comandos completos", level=1)
add_code(
    doc,
    r'cd "C:\Usil\6to_ciclo_Compiladores\Trabajo final_compiladores"' + "\n"
    "git status\n"
    "git add .\n"
    'git commit -m "Configura despliegue en Railway con Docker"\n'
    "git pull --rebase origin main\n"
    "git push -u origin main"
)

doc.add_heading("Resultado esperado", level=1)
doc.add_paragraph(
    "Al finalizar, Railway mostrara una URL publica. Al abrirla, deberias ver la interfaz web del DSL. "
    "Cuando presiones Ejecutar workflow, Flask llamara al compilador C++ compilado en Linux y mostrara tokens, logs y errores."
)

doc.save(OUT)
print(OUT)
