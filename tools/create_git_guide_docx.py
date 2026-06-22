from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Guia_control_versiones_github.docx"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(table, top=80, start=120, bottom=80, end=120):
    tbl_pr = table._tbl.tblPr
    margins = tbl_pr.first_child_found_in("w:tblCellMar")
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        tbl_pr.append(margins)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(34, 43, 54)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.add_run(text)


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.add_run(text)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.10

for name, size, color, before, after in [
    ("Heading 1", 16, "2E74B5", 16, 8),
    ("Heading 2", 13, "2E74B5", 12, 6),
    ("Heading 3", 12, "1F4D78", 8, 4),
]:
    style = styles[name]
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = title.add_run("Guia para controlar versiones con Git y GitHub")
r.font.name = "Calibri"
r.font.size = Pt(24)
r.font.bold = True
r.font.color.rgb = RGBColor(11, 37, 69)

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(12)
subtitle.add_run("Proyecto: Trabajo final_compiladores").bold = True

doc.add_heading("Que significa el error", level=1)
doc.add_paragraph(
    "El mensaje 'rejected main -> main (fetch first)' significa que el repositorio de GitHub ya tiene commits "
    "en la rama main que tu carpeta local no tiene. Git bloquea el push para evitar que borres o sobreescribas "
    "trabajo anterior sin revisarlo."
)

doc.add_heading("Ruta recomendada: conservar versiones", level=1)
doc.add_paragraph("Usa esta ruta si quieres mantener el historial anterior del repositorio y subir este proyecto correctamente.")
add_number(doc, "Verifica que estas en la carpeta del proyecto.")
add_code(doc, r'cd "C:\Usil\6to_ciclo_Compiladores\Trabajo final_compiladores"')
add_number(doc, "Revisa el estado y el remoto.")
add_code(doc, "git status\ngit remote -v")
add_number(doc, "Crea una rama de respaldo local antes de mezclar historias.")
add_code(doc, 'git branch respaldo-trabajo-final')
add_number(doc, "Trae la informacion actual de GitHub.")
add_code(doc, "git fetch origin")
add_number(doc, "Integra los cambios remotos con tu trabajo local.")
add_code(doc, "git pull --rebase origin main")
add_number(doc, "Si no hay conflictos, sube el proyecto.")
add_code(doc, "git push -u origin main")

doc.add_heading("Si aparece 'unrelated histories'", level=1)
doc.add_paragraph(
    "Ese mensaje aparece cuando el repositorio remoto y tu carpeta local fueron creados por separado. En ese caso, usa:"
)
add_code(doc, "git pull origin main --allow-unrelated-histories")
doc.add_paragraph("Luego resuelve conflictos si Git los muestra, confirma la fusion y sube:")
add_code(doc, "git add .\ngit commit -m \"Integra proyecto final con historial remoto\"\ngit push -u origin main")

doc.add_heading("Como resolver conflictos", level=1)
add_bullet(doc, "Git marcara archivos con conflictos. Abrelos y busca marcas como <<<<<<<, ======= y >>>>>>>.")
add_bullet(doc, "Elige que contenido conservar: el anterior de GitHub, el nuevo del proyecto, o una combinacion.")
add_bullet(doc, "Despues de editar, guarda los archivos y ejecuta git add .")
add_bullet(doc, "Si estabas usando rebase, ejecuta git rebase --continue. Si estabas usando merge, ejecuta git commit.")

doc.add_heading("Control de versiones recomendado", level=1)
doc.add_paragraph("Para trabajar ordenadamente, usa commits, ramas y etiquetas.")

table = doc.add_table(rows=1, cols=3)
table.style = "Table Grid"
set_cell_margins(table)
set_table_width(table, [1.4, 2.2, 2.9])
hdr = table.rows[0].cells
hdr[0].text = "Elemento"
hdr[1].text = "Para que sirve"
hdr[2].text = "Ejemplo"
for cell in hdr:
    shade(cell, "F2F4F7")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

rows = [
    ("Commit", "Guarda un cambio puntual en el historial.", 'git commit -m "Actualiza guia de instalacion"'),
    ("Rama", "Permite trabajar una version sin tocar main.", "git checkout -b mejora-interfaz"),
    ("Tag", "Marca una entrega estable o version final.", 'git tag -a v1.0 -m "Entrega final"'),
]
for a, b, c in rows:
    cells = table.add_row().cells
    cells[0].text = a
    cells[1].text = b
    cells[2].text = c

doc.add_heading("Flujo practico para tus entregas", level=1)
add_code(
    doc,
    "git status\n"
    "git add .\n"
    'git commit -m "Describe claramente el cambio"\n'
    "git pull --rebase origin main\n"
    "git push origin main"
)

doc.add_heading("Crear una version estable", level=1)
doc.add_paragraph("Cuando el proyecto ya corre bien, crea una etiqueta de version:")
add_code(doc, 'git tag -a v1.0 -m "Entrega final compiladores"\ngit push origin v1.0')

doc.add_heading("Opcion no recomendada: reemplazar lo anterior", level=1)
doc.add_paragraph(
    "Solo usa esta opcion si estas completamente seguro de que el contenido anterior del repositorio no importa. "
    "Este comando reescribe la rama main remota."
)
add_code(doc, "git push --force-with-lease origin main")
doc.add_paragraph(
    "No uses git push --force como primera opcion. --force-with-lease es menos riesgoso, pero igualmente puede cambiar "
    "el historial remoto."
)

doc.add_heading("Comandos utiles", level=1)
add_code(
    doc,
    "git status\n"
    "git log --oneline --decorate --graph --all\n"
    "git remote -v\n"
    "git branch\n"
    "git tag\n"
    "git diff\n"
    "git restore --staged archivo"
)

doc.add_heading("Recomendacion final", level=1)
doc.add_paragraph(
    "Para tu caso, la opcion mas segura es: crear respaldo, hacer git fetch, luego git pull --rebase origin main, "
    "resolver conflictos si aparecen y finalmente ejecutar git push -u origin main. Asi conservas el historial anterior "
    "y agregas esta entrega como una nueva version del repositorio."
)

doc.save(OUT)
print(OUT)
