from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Guia_instalacion_flask_y_uso_web.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(table, top=80, start=120, bottom=80, end=120):
    tbl_pr = table._tbl.tblPr
    margins = tbl_pr.first_child_found_in("w:tblCellMar")
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        tbl_pr.append(margins)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(34, 43, 54)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.add_run(text)


def add_step(doc, title, body, command=None):
    h = doc.add_paragraph(style="Heading 2")
    h.add_run(title)
    p = doc.add_paragraph(body)
    if command:
        add_code(doc, command)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.10

for name, size, color, before, after in [
    ("Heading 1", 16, "2E74B5", 16, 8),
    ("Heading 2", 13, "2E74B5", 12, 6),
    ("Heading 3", 12, "1F4D78", 8, 4),
]:
    style = styles[name]
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.paragraph_format.space_after = Pt(3)
r = title.add_run("Guia de instalacion y uso web")
r.font.name = "Calibri"
r.font.size = Pt(24)
r.font.bold = True
r.font.color.rgb = RGBColor(11, 37, 69)

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(12)
subtitle.add_run("Proyecto: DSL Delivery Workflow Compiler con Flask y C++").bold = True

doc.add_heading("Objetivo de esta guia", level=1)
doc.add_paragraph(
    "Este documento explica como preparar el entorno, instalar Flask, compilar el compilador en C++ "
    "y abrir la aplicacion web desde un navegador."
)

doc.add_heading("Requisitos", level=1)
add_bullet(doc, "Python 3 instalado y disponible desde la terminal.")
add_bullet(doc, "Un compilador C++ compatible con C++17: MinGW-w64, LLVM/Clang o Visual Studio Build Tools.")
add_bullet(doc, "Acceso a la carpeta del proyecto en Windows.")
add_bullet(doc, "Navegador web: Chrome, Edge, Firefox u otro equivalente.")

doc.add_heading("Estructura del proyecto", level=1)
table = doc.add_table(rows=1, cols=2)
table.style = "Table Grid"
set_cell_margins(table)
set_table_width(table, [2.0, 4.5])
hdr = table.rows[0].cells
hdr[0].text = "Carpeta o archivo"
hdr[1].text = "Proposito"
for cell in hdr:
    set_cell_shading(cell, "F2F4F7")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

rows = [
    ("compiler/main.cpp", "Codigo fuente del compilador DSL en C++."),
    ("web/app.py", "Servidor Flask que conecta la pagina web con el compilador."),
    ("web/templates", "Vista HTML principal de la aplicacion."),
    ("web/static", "CSS y JavaScript de la interfaz."),
    ("examples", "Programas DSL de prueba."),
    ("requirements.txt", "Dependencias de Python necesarias para ejecutar Flask."),
    ("build.ps1", "Script para compilar el programa C++ en Windows."),
]
for left, right in rows:
    cells = table.add_row().cells
    cells[0].text = left
    cells[1].text = right
    for cell in cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

doc.add_heading("Instalacion paso a paso", level=1)
add_step(
    doc,
    "1. Abrir PowerShell en la carpeta del proyecto",
    "Ubicate en la carpeta donde se encuentran README.md, build.ps1 y requirements.txt.",
    r"cd C:\Users\delar\Documents\Codex\2026-05-24\tu-proyecto-combina-tres-reas-importantes",
)
add_step(
    doc,
    "2. Crear un entorno virtual",
    "Este paso guarda las dependencias dentro del proyecto y evita errores de permisos en Windows.",
    "python -m venv .venv",
)
add_step(
    doc,
    "3. Instalar Flask dentro del entorno virtual",
    "Ejecuta el siguiente comando para instalar las dependencias declaradas en requirements.txt.",
    r".\.venv\Scripts\python.exe -m pip install -r requirements.txt",
)
add_step(
    doc,
    "4. Compilar el compilador C++",
    "Este comando genera el ejecutable bin\\delivery_compiler.exe. La web lo necesita para analizar el DSL.",
    r".\build.ps1",
)
add_step(
    doc,
    "5. Ejecutar la aplicacion Flask",
    "Cuando el servidor este activo, la terminal mostrara una direccion local.",
    r".\.venv\Scripts\python.exe web\app.py",
)
add_step(
    doc,
    "6. Abrir la pagina web",
    "Abre el navegador y entra a la siguiente direccion local.",
    "http://127.0.0.1:5000",
)

doc.add_heading("Como usar la pagina", level=1)
add_bullet(doc, "Escribe o modifica el codigo DSL en el editor de la izquierda.")
add_bullet(doc, "Presiona Ejecutar workflow.")
add_bullet(doc, "Revisa el estado, los logs, errores y tokens reconocidos.")
add_bullet(doc, "Si aparece un error indicando que falta bin/delivery_compiler.exe, vuelve a ejecutar .\\build.ps1.")

doc.add_heading("Ejemplo de DSL valido", level=1)
add_code(
    doc,
    'PEDIDO {\n'
    '    cliente: "Carlos"\n'
    '    producto: "Pizza Familiar"\n'
    '    total: 80\n'
    '    pago: YAPE\n'
    '    direccion: "Av. Lima 123"\n'
    '    stock: 4\n'
    '}\n\n'
    'VALIDAR stock\n'
    'VALIDAR direccion\n'
    'VALIDAR pago\n\n'
    'SI total > 50 {\n'
    '    ASIGNAR prioridad_alta\n'
    '}\n\n'
    'ASIGNAR repartidor\n'
    'INICIAR entrega\n'
    'FINALIZAR pedido'
)

doc.add_heading("Errores comunes", level=1)
error_table = doc.add_table(rows=1, cols=2)
error_table.style = "Table Grid"
set_cell_margins(error_table)
set_table_width(error_table, [2.4, 4.1])
error_hdr = error_table.rows[0].cells
error_hdr[0].text = "Problema"
error_hdr[1].text = "Solucion"
for cell in error_hdr:
    set_cell_shading(cell, "F2F4F7")

for problem, solution in [
    ("Flask no esta instalado", "Ejecuta .\\.venv\\Scripts\\python.exe -m pip install -r requirements.txt."),
    ("Error de permisos al instalar paquetes", "Usa el entorno virtual: python -m venv .venv."),
    ("No existe delivery_compiler.exe", "Ejecuta .\\build.ps1 desde la raiz del proyecto."),
    ("No se encontro compilador C++", "Instala MinGW-w64, LLVM/Clang o Visual Studio Build Tools."),
    ("El puerto 5000 esta ocupado", "Cierra el otro servidor o cambia el puerto en web/app.py."),
    ("Error semantico en VALIDAR stock", "Asegurate de declarar stock dentro del bloque PEDIDO y que sea mayor que cero."),
]:
    cells = error_table.add_row().cells
    cells[0].text = problem
    cells[1].text = solution

doc.add_heading("Comandos principales", level=1)
add_code(
    doc,
    "python -m venv .venv\n"
    ".\\.venv\\Scripts\\python.exe -m pip install -r requirements.txt\n"
    ".\\build.ps1\n"
    ".\\.venv\\Scripts\\python.exe web\\app.py\n"
    "http://127.0.0.1:5000"
)

doc.add_heading("Cierre", level=1)
doc.add_paragraph(
    "Con estos pasos tendras una aplicacion web funcional que envia codigo DSL a un compilador C++, "
    "recibe una salida JSON y muestra el resultado del workflow logistico en el navegador."
)

doc.save(OUT)
print(OUT)
